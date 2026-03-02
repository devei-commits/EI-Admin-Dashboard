import subprocess
from datetime import datetime, timedelta
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


def main() -> None:
    root = Path(r"H:/Eisthetic/eiadmin/eiadmin")
    repos = {
        "eiadmin": root,
        "Esthetic-Insights-Website": root / "Esthetic-Insights-Website",
    }

    # Assumption based on user request: work period from Jan 1, 2026 to today.
    start_date = datetime(2026, 1, 1).date()
    end_date = datetime(2026, 3, 2).date()

    all_days = []
    current = start_date
    while current <= end_date:
        all_days.append(current)
        current += timedelta(days=1)

    per_day_commits: dict[datetime.date, list[dict]] = defaultdict(list)

    for repo_name, repo_path in repos.items():
        if not repo_path.exists():
            continue

        try:
            out = subprocess.check_output(
                [
                    "git",
                    "log",
                    f"--since={start_date.isoformat()}",
                    f"--until={end_date.isoformat()} 23:59:59",
                    "--date=short",
                    "--pretty=%ad\t%H\t%an\t%s",
                ],
                cwd=str(repo_path),
                text=True,
                errors="ignore",
            )
        except subprocess.CalledProcessError:
            continue

        for line in out.splitlines():
            if not line.strip():
                continue
            parts = line.split("\t", 3)
            if len(parts) != 4:
                continue
            date_str, commit_hash, author, subject = parts
            try:
                day = datetime.strptime(date_str, "%Y-%m-%d").date()
            except ValueError:
                continue
            per_day_commits[day].append(
                {
                    "repo": repo_name,
                    "hash": commit_hash,
                    "author": author,
                    "subject": subject.strip(),
                }
            )

    wb = Workbook()
    ws = wb.active
    ws.title = "Work Log"

    ws.append(["Date", "Work description", "Output"])

    for day in all_days:
        commits = per_day_commits.get(day, [])
        date_str_disp = day.strftime("%Y-%m-%d")

        if not commits:
            work_desc = "No git commits recorded."
            output_desc = "No code output (no commits)."
        else:
            by_repo: dict[str, list[dict]] = defaultdict(list)
            for c in commits:
                by_repo[c["repo"]].append(c)

            desc_lines: list[str] = []
            output_parts: list[str] = []
            for repo_name, repo_commits in sorted(by_repo.items()):
                subjects = [c["subject"] for c in repo_commits]
                unique_subjects: list[str] = []
                seen = set()
                for s in subjects:
                    if s not in seen:
                        seen.add(s)
                        unique_subjects.append(s)
                if unique_subjects:
                    desc_lines.append(f"{repo_name}: " + "; ".join(unique_subjects))
                output_parts.append(f"{repo_name}: {len(repo_commits)} commits")

            work_desc = "\n".join(desc_lines)
            total_commits = len(commits)
            output_desc = (
                f"Total commits: {total_commits} (" + ", ".join(output_parts) + ")"
            )

        ws.append([date_str_disp, work_desc, output_desc])

    for col_idx, width in [(1, 12), (2, 80), (3, 40)]:
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    out_path = root / "EI_Work_Log_From_2026-01-01_to_2026-03-02.xlsx"
    wb.save(out_path)
    print(f"Worksheet written to: {out_path}")


if __name__ == "__main__":
    main()

